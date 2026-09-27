"""
Renders MonthlyReportData into a PDF (reportlab) and an Excel workbook
(openpyxl), per spec section 14. Both are written to disk and the paths
returned, so the caller (API endpoint or scheduled job) can attach/email them.
"""

import os

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import openpyxl
from openpyxl.styles import Font, PatternFill

from app.reports.report_engine import MonthlyReportData

REPORT_DIR = "/tmp/pm_reports"
os.makedirs(REPORT_DIR, exist_ok=True)


def render_pdf(report: MonthlyReportData) -> str:
    path = os.path.join(REPORT_DIR, f"pm_report_{report.month}_{report.financial_year}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Preventive Maintenance Performance Report", styles["Title"]))
    elements.append(Paragraph(f"Month: {report.month} {report.financial_year}", styles["Normal"]))
    elements.append(Spacer(1, 10))

    summary_data = [
        ["Total PM Planned", report.total_planned],
        ["PM Completed", report.completed],
        ["Pending", report.pending],
        ["Overdue", report.overdue],
        ["Completion Rate", f"{report.completion_rate}%"],
        ["On-Time Completion", f"{report.on_time_rate}%"],
        ["Late Completion", f"{report.late_rate}%"],
    ]
    elements.append(_styled_table([["Metric", "Value"]] + summary_data))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Machine-wise Report", styles["Heading2"]))
    rows = [["Machine", "Plan", "Actual", "Pending", "Overdue", "Perf %"]]
    for r in report.machine_rows:
        rows.append([f"{r['machine_number']} - {r['machine_name']}", r["plan"], r["actual"],
                     r["pending"], r["overdue"], f"{r['performance_pct']}%"])
    elements.append(_styled_table(rows))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Location-wise Report", styles["Heading2"]))
    rows = [["Location", "Plan", "Actual", "Pending", "Perf %"]]
    for r in report.location_rows:
        rows.append([r["location"], r["plan"], r["actual"], r["pending"], f"{r['performance_pct']}%"])
    elements.append(_styled_table(rows))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Overdue PM List", styles["Heading2"]))
    rows = [["Machine", "Planned Date", "Days Overdue", "Assigned To"]]
    for r in report.overdue_list:
        rows.append([f"{r['machine_number']} - {r['machine_name']}", r["planned_date"],
                     r["days_overdue"], r["assigned_to"]])
    elements.append(_styled_table(rows) if len(rows) > 1 else Paragraph("None.", styles["Normal"]))

    doc.build(elements)
    return path


def _styled_table(rows) -> Table:
    t = Table(rows, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1d21")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dfe3e8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f5f7")]),
    ]))
    return t


def render_excel(report: MonthlyReportData) -> str:
    path = os.path.join(REPORT_DIR, f"pm_report_{report.month}_{report.financial_year}.xlsx")
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Summary"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A1D21")

    ws.append(["Preventive Maintenance Performance Report"])
    ws.append([f"Month: {report.month} {report.financial_year}"])
    ws.append([])
    ws.append(["Metric", "Value"])
    for cell in ws[4]:
        cell.font = header_font
        cell.fill = header_fill
    for row in [
        ("Total PM Planned", report.total_planned),
        ("PM Completed", report.completed),
        ("Pending", report.pending),
        ("Overdue", report.overdue),
        ("Completion Rate %", report.completion_rate),
        ("On-Time Completion %", report.on_time_rate),
        ("Late Completion %", report.late_rate),
    ]:
        ws.append(row)

    _write_sheet(wb, "Machine-wise", ["Machine No", "Machine Name", "Plan", "Actual", "Pending", "Overdue", "Perf %"],
                 [[r["machine_number"], r["machine_name"], r["plan"], r["actual"], r["pending"],
                   r["overdue"], r["performance_pct"]] for r in report.machine_rows])

    _write_sheet(wb, "Location-wise", ["Location", "Plan", "Actual", "Pending", "Perf %"],
                 [[r["location"], r["plan"], r["actual"], r["pending"], r["performance_pct"]]
                  for r in report.location_rows])

    _write_sheet(wb, "Employee-wise", ["Employee", "Assigned", "Completed", "Pending", "Overdue"],
                 [[r["employee"], r["assigned"], r["completed"], r["pending"], r["overdue"]]
                  for r in report.employee_rows])

    _write_sheet(wb, "Overdue PM", ["Machine No", "Machine Name", "Planned Date", "Days Overdue", "Assigned To"],
                 [[r["machine_number"], r["machine_name"], r["planned_date"], r["days_overdue"], r["assigned_to"]]
                  for r in report.overdue_list])

    wb.save(path)
    return path


def _write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A1D21")
    for row in rows:
        ws.append(row)
