"""
Bulk import from Excel for master data that is otherwise typed in by hand:
spare parts and employees.

Flow (stateless - the browser re-sends the same file for each step):
  GET  /api/bulk-import/{kind}/template  -> downloadable .xlsx with the right columns
  POST /api/bulk-import/{kind}/preview   -> validates + diffs against the database, writes NOTHING
  POST /api/bulk-import/{kind}/commit    -> applies every valid row in ONE transaction;
                                            rows with errors are skipped and reported

Rules that keep data safe:
  * Spare parts: stock of EXISTING parts is never changed by an import (stock only moves
    through the transaction ledger). Opening stock applies to NEW parts only and is
    recorded as an "Initial stock" ledger entry, exactly like adding a part by hand.
  * Employees: passwords of EXISTING people are never touched. New people get the
    password from the file, or a generated one that is shown once after the import.
  * Blank cells never overwrite existing values.
"""
import io
import re
import secrets
import string

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.audit import record_audit
from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.security import hash_password
from app.models.models import Employee, Role, SparePart, SparePartTransaction

router = APIRouter(prefix="/api/bulk-import", tags=["bulk-import"])

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_ROWS_IN_RESPONSE = 500
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

KIND_ROLES = {
    "spare-parts": {Role.SUPERVISOR, Role.MANAGER, Role.ADMIN},
    "employees": {Role.ADMIN},
}

# field -> (header shown in template, required, accepted header spellings)
SPARE_PART_COLUMNS = {
    "part_code": ("Part Code *", True, ["part code", "code", "part no", "part number", "item code", "sku"]),
    "name": ("Name *", True, ["name", "part name", "item name", "item"]),
    "description": ("Description", False, ["description", "desc", "remarks", "details"]),
    "unit": ("Unit", False, ["unit", "uom", "unit of measure"]),
    "opening_stock": ("Opening Stock", False, ["opening stock", "stock", "stock on hand", "qty", "quantity", "initial stock"]),
    "minimum_stock": ("Minimum Stock", False, ["minimum stock", "min stock", "min qty", "reorder level", "reorder point"]),
    "unit_cost": ("Unit Cost", False, ["unit cost", "cost", "price", "rate"]),
    "storage_location": ("Storage Location", False, ["storage location", "location", "bin", "rack", "store"]),
    "preferred_vendor": ("Preferred Vendor", False, ["preferred vendor", "vendor", "supplier"]),
}
EMPLOYEE_COLUMNS = {
    "name": ("Name *", True, ["name", "employee name", "full name"]),
    "email": ("Email *", True, ["email", "email address", "e mail", "mail id"]),
    "role": ("Role", False, ["role", "access", "access level"]),
    "department": ("Department", False, ["department", "dept"]),
    "designation": ("Designation", False, ["designation", "title", "job title", "position"]),
    "phone": ("Phone", False, ["phone", "mobile", "phone number", "mobile number", "whatsapp"]),
    "password": ("Password", False, ["password", "temporary password", "temp password"]),
}
COLUMNS = {"spare-parts": SPARE_PART_COLUMNS, "employees": EMPLOYEE_COLUMNS}

INSTRUCTIONS = {
    "spare-parts": [
        "Fill the 'Data' sheet, one spare part per row. Columns marked * are required.",
        "Part Code is the unique ID. If it already exists, the part is UPDATED (blank cells are left as they are).",
        "Opening Stock only applies to NEW parts. Stock of existing parts is never changed by an import - use stock adjustments for that.",
        "Unit: pcs, ltr, kg, mtr ... (default pcs). Stock and Minimum Stock must be whole numbers. Unit Cost is in whole currency units.",
        "Example row:  BRG-6204 | Ball bearing 6204 | Deep groove | pcs | 25 | 5 | 180 | Rack A2 | SKF Dealer",
    ],
    "employees": [
        "Fill the 'Data' sheet, one person per row. Columns marked * are required.",
        "Email is the unique ID. If it already exists, the person is UPDATED (blank cells are left as they are).",
        "Role: ADMIN, MANAGER, SUPERVISOR, TECHNICIAN or VIEWER (default TECHNICIAN).",
        "Password (new people only, min 8 characters). Leave blank to generate one - it is shown once after the import.",
        "Existing people's passwords are never changed by an import.",
        "Example row:  Ravi Kumar | ravi@company.com | TECHNICIAN | Maintenance | Fitter | 9876543210 |",
    ],
}


# ---------------------------------------------------------------- helpers
def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(h or "").lower()).strip()


def _s(v):
    """Cell -> trimmed string or None. Whole-number floats (1001.0) become '1001'."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _whole(v, label: str):
    """Cell -> non-negative int, or raises ValueError with a readable message."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        f = float(str(v).replace(",", "").strip())
    except ValueError:
        raise ValueError(f"{label} must be a number")
    if f < 0:
        raise ValueError(f"{label} cannot be negative")
    if not f.is_integer():
        raise ValueError(f"{label} must be a whole number")
    return int(f)


def _money(v, label: str):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        f = float(str(v).replace(",", "").strip())
    except ValueError:
        raise ValueError(f"{label} must be a number")
    if f < 0:
        raise ValueError(f"{label} cannot be negative")
    return int(round(f))


def _check_access(kind: str, user: Employee):
    if kind not in COLUMNS:
        raise HTTPException(404, "Unknown import type")
    if user.role not in KIND_ROLES[kind]:
        raise HTTPException(403, "Insufficient permissions for this import")


def _read_rows(kind: str, file: UploadFile):
    """Parse the workbook into [(row_number, {field: raw_value})]. Raises HTTPException on a bad file."""
    data = file.file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(400, "File is larger than 5 MB")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "Could not read the file. Please upload a .xlsx Excel file (use the template).")
    ws = wb.worksheets[0]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise HTTPException(400, "The first sheet is empty")

    spec = COLUMNS[kind]
    alias_to_field = {}
    for field, (_label, _req, aliases) in spec.items():
        for a in aliases:
            alias_to_field[a] = field
    col_index = {}
    for i, h in enumerate(header):
        field = alias_to_field.get(_norm(h))
        if field and field not in col_index:
            col_index[field] = i
    missing = [spec[f][0].replace(" *", "") for f, (_l, req, _a) in spec.items() if req and f not in col_index]
    if missing:
        raise HTTPException(400, f"Missing required column(s): {', '.join(missing)}. Download the template to see the expected columns.")

    out = []
    for n, row in enumerate(rows_iter, start=2):
        values = {f: (row[i] if i < len(row) else None) for f, i in col_index.items()}
        if all(_s(v) is None for v in values.values()):
            continue  # blank line
        out.append((n, values))
        if len(out) > MAX_ROWS:
            raise HTTPException(400, f"Too many rows (limit {MAX_ROWS}). Split the file and import in parts.")
    if not out:
        raise HTTPException(400, "No data rows found under the header row")
    return out


# ---------------------------------------------------------------- spare parts
def _plan_spare_parts(db: Session, raw_rows):
    codes = {_s(v.get("part_code")) for _n, v in raw_rows if _s(v.get("part_code"))}
    existing = {p.part_code: p for p in db.query(SparePart).filter(SparePart.part_code.in_(codes)).all()} if codes else {}
    seen = {}
    plan = []
    for n, v in raw_rows:
        errors, warnings = [], []
        code, name = _s(v.get("part_code")), _s(v.get("name"))
        if not code:
            errors.append("Part Code is required")
        if not name:
            errors.append("Name is required")
        if code and len(code) > 50:
            errors.append("Part Code is longer than 50 characters")
        if code in seen:
            errors.append(f"Duplicate Part Code (first used in row {seen[code]})")
        elif code:
            seen[code] = n

        fields = {}
        for key, conv, label in (("opening_stock", _whole, "Opening Stock"), ("minimum_stock", _whole, "Minimum Stock"),
                                 ("unit_cost", _money, "Unit Cost")):
            try:
                fields[key] = conv(v.get(key), label)
            except ValueError as e:
                errors.append(str(e))
        for key in ("description", "unit", "storage_location", "preferred_vendor"):
            fields[key] = _s(v.get(key))
        if name:
            fields["name"] = name

        action = "error"
        part = existing.get(code) if code else None
        if not errors:
            if part is None:
                action = "create"
            else:
                changed = {k: val for k, val in fields.items()
                           if k not in ("opening_stock",) and val is not None and getattr(part, k) != val}
                if fields.get("opening_stock") not in (None, 0):
                    warnings.append("Opening Stock ignored for an existing part (use stock adjustment)")
                action = "update" if changed else "unchanged"
                fields = changed
        plan.append({"row": n, "key": code or "", "name": name or "", "action": action,
                     "messages": errors + warnings, "fields": fields, "part": part})
    return plan


def _apply_spare_parts(db: Session, plan, user: Employee):
    created = updated = 0
    for item in plan:
        if item["action"] == "create":
            f = item["fields"]
            stock = f.get("opening_stock") or 0
            part = SparePart(
                part_code=item["key"], name=f["name"], description=f.get("description"),
                unit=f.get("unit") or "pcs", minimum_stock=f.get("minimum_stock") or 0,
                stock_on_hand=stock, unit_cost=f.get("unit_cost"),
                storage_location=f.get("storage_location"), preferred_vendor=f.get("preferred_vendor"),
            )
            db.add(part)
            db.flush()
            if stock:
                db.add(SparePartTransaction(spare_part_id=part.id, change=stock,
                                            reason="Initial stock (Excel import)", performed_by=user.id))
            created += 1
        elif item["action"] == "update":
            for k, val in item["fields"].items():
                setattr(item["part"], k, val)
            updated += 1
    return {"created": created, "updated": updated}


# ---------------------------------------------------------------- employees
_PW_ALPHABET = "abcdefghijkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"


def _gen_password() -> str:
    return "".join(secrets.choice(_PW_ALPHABET) for _ in range(10))


def _plan_employees(db: Session, raw_rows, user: Employee):
    emails = {(_s(v.get("email")) or "").lower() for _n, v in raw_rows if _s(v.get("email"))}
    existing = {e.email.lower(): e for e in db.query(Employee).filter(func.lower(Employee.email).in_(emails)).all() if e.email} if emails else {}
    seen = {}
    plan = []
    for n, v in raw_rows:
        errors, warnings = [], []
        name, email = _s(v.get("name")), (_s(v.get("email")) or "").lower()
        if not name:
            errors.append("Name is required")
        if not email:
            errors.append("Email is required")
        elif not EMAIL_RE.match(email):
            errors.append("Email is not valid")
        elif email in seen:
            errors.append(f"Duplicate email (first used in row {seen[email]})")
        else:
            seen[email] = n

        role = None
        role_raw = _s(v.get("role"))
        if role_raw:
            try:
                role = Role[role_raw.strip().upper()]
            except KeyError:
                errors.append(f"Role '{role_raw}' is not valid (use {', '.join(r.name for r in Role)})")
        password = _s(v.get("password"))
        if password and len(password) < 8:
            errors.append("Password must be at least 8 characters")

        fields = {k: _s(v.get(k)) for k in ("department", "designation", "phone")}
        if name:
            fields["name"] = name
        emp = existing.get(email) if email else None
        action = "error"
        if not errors:
            if emp is None:
                action = "create"
                fields["role"] = role or Role.TECHNICIAN
                fields["password"] = password
            else:
                if role is not None:
                    if emp.id == user.id and role != emp.role:
                        warnings.append("Your own role was not changed")
                    else:
                        fields["role"] = role
                if password:
                    warnings.append("Password ignored for an existing person")
                changed = {k: val for k, val in fields.items() if val is not None and getattr(emp, k) != val}
                action = "update" if changed else "unchanged"
                fields = changed
        plan.append({"row": n, "key": email, "name": name or "", "action": action,
                     "messages": errors + warnings, "fields": fields, "emp": emp})
    return plan


def _apply_employees(db: Session, plan):
    created = updated = 0
    generated = []
    for item in plan:
        if item["action"] == "create":
            f = item["fields"]
            pw = f.get("password")
            generated_pw = None
            if not pw:
                pw = generated_pw = _gen_password()
            db.add(Employee(name=f["name"], email=item["key"], hashed_password=hash_password(pw),
                            role=f["role"], department=f.get("department"),
                            designation=f.get("designation"), phone=f.get("phone")))
            created += 1
            if generated_pw:
                generated.append({"name": f["name"], "email": item["key"], "password": generated_pw})
        elif item["action"] == "update":
            for k, val in item["fields"].items():
                setattr(item["emp"], k, val)
            updated += 1
    return {"created": created, "updated": updated}, generated


# ---------------------------------------------------------------- endpoints
def _summarize(kind: str, plan):
    totals = {"rows": len(plan), "create": 0, "update": 0, "unchanged": 0, "error": 0}
    for i in plan:
        totals[i["action"]] += 1
    rows = [{"row": i["row"], "key": i["key"], "name": i["name"], "action": i["action"], "messages": i["messages"]}
            for i in plan]
    # show problems first so they are never hidden by the response cap
    rows.sort(key=lambda r: (r["action"] != "error", r["row"]))
    return {"kind": kind, "totals": totals, "rows": rows[:MAX_ROWS_IN_RESPONSE],
            "truncated": len(rows) > MAX_ROWS_IN_RESPONSE}


def _build_plan(kind: str, db: Session, file: UploadFile, user: Employee):
    raw = _read_rows(kind, file)
    return _plan_spare_parts(db, raw) if kind == "spare-parts" else _plan_employees(db, raw, user)


@router.get("/{kind}/template")
def download_template(kind: str, user: Employee = Depends(get_current_user)):
    _check_access(kind, user)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    fill = PatternFill("solid", fgColor="DDE7F5")
    for col, (label, _req, _a) in enumerate(COLUMNS[kind].values(), start=1):
        c = ws.cell(row=1, column=col, value=label)
        c.font = Font(bold=True)
        c.fill = fill
        ws.column_dimensions[c.column_letter].width = max(16, len(label) + 4)
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Instructions")
    info.column_dimensions["A"].width = 120
    for i, line in enumerate(INSTRUCTIONS[kind], start=1):
        info.cell(row=i, column=1, value=line)
    buf = io.BytesIO()
    wb.save(buf)
    filename = f"{kind}-import-template.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/{kind}/preview")
def preview_import(kind: str, file: UploadFile = File(...), db: Session = Depends(get_db),
                   user: Employee = Depends(get_current_user)):
    _check_access(kind, user)
    return _summarize(kind, _build_plan(kind, db, file, user))


@router.post("/{kind}/commit")
def commit_import(kind: str, file: UploadFile = File(...), db: Session = Depends(get_db),
                  user: Employee = Depends(get_current_user)):
    _check_access(kind, user)
    plan = _build_plan(kind, db, file, user)
    result = _summarize(kind, plan)

    generated = []
    try:
        if kind == "spare-parts":
            applied = _apply_spare_parts(db, plan, user)
        else:
            applied, generated = _apply_employees(db, plan)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Someone else added one of these records at the same time. Please preview and try again.")

    record_audit(
        db, action="BULK_IMPORT_SPARE_PARTS" if kind == "spare-parts" else "BULK_IMPORT_EMPLOYEES",
        entity_type="BulkImport", entity_id=None, actor_id=user.id,
        new_value={**applied, "skipped_errors": result["totals"]["error"],
                   "created_keys": [i["key"] for i in plan if i["action"] == "create"][:100]},
    )
    result["applied"] = applied
    if generated:
        result["generated_passwords"] = generated
    return result
